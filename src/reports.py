import json

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .config import ARTIFACTS, FIGURES, REPORTS, SCORED_PATH


def _base():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    return doc


def _table(doc, frame, widths=None):
    t = doc.add_table(rows=1, cols=len(frame.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(frame.columns):
        t.rows[0].cells[i].text = str(c)
        for p in t.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
    for _, row in frame.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()
    return t


def _fig(doc, name, caption, width=6.0):
    doc.add_picture(str(FIGURES / name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def _facts():
    d = pd.read_csv(SCORED_PATH)
    m = json.loads((ARTIFACTS / "metrics.json").read_text())
    summ = pd.read_csv(ARTIFACTS / "cluster_summary.csv", index_col=0)
    lev = pd.read_csv(ARTIFACTS / "lever_importance.csv", index_col=0)["EtaSquared"]
    k = pd.read_csv(ARTIFACTS / "k_selection.csv")
    ch = {}
    for c in ["InStore", "UberEats", "DoorDash", "SelfDelivery"]:
        ch[c] = (
            d[f"{c}NetProfit"].sum() / d[f"{c}Revenue"].sum() * 100,
            (d[f"{c}NetProfit"] < 0).mean() * 100,
        )
    seg = d.groupby("Segment")[["NetMargin", "CostBurden", "AggregatorDependence", "GPI"]].mean()
    sub = d.groupby("Subregion")[["GrowthFactor", "GPI", "NetMargin"]].mean()
    prof = d.groupby("ClusterLabel")["TotalNetProfit"].sum()
    return d, m, summ, lev, k, ch, seg, sub, prof


def build_paper():
    d, m, summ, lev, k, ch, seg, sub, prof = _facts()
    doc = _base()
    doc.add_heading("Restaurant Growth Potential Modeling & Strategic Classification System", 0)
    doc.add_paragraph("SkyCity Auckland Restaurants & Bars | Unified Mentor Internship Program | Research Paper")
    doc.add_paragraph("Author: Rimsha Naaz")

    doc.add_heading("Abstract", 1)
    top = summ.index[0]
    doc.add_paragraph(
        f"Expansion decisions across {m['n_restaurants']:,} SkyCity Auckland restaurant branches are usually uniform and intuition-based. "
        f"This study builds a growth-readiness lens from {len(pd.read_csv(ARTIFACTS / 'cluster_centroids_zscore.csv', index_col=0).columns)} engineered "
        f"features covering demand, channel mix, cost structure and logistics. After skew correction and scaling, PCA compressed the data to "
        f"{m['pca_components']} components ({m['pca_cumulative_used'] * 100:.1f}% of variance). K-Means with k = {m['chosen_k']} produced six interpretable archetypes, "
        f"validated against hierarchical clustering (ARI {m['robustness']['kmeans_vs_hierarchical_ARI']:.2f}). A transparent Growth Potential Index (GPI, 0-100) scores every branch, "
        f"and a rule layer assigns one of three actions: Optimize, Rebalance channels, or Hold / Stabilize. "
        f"The strongest archetype is {top} (mean GPI {summ.iloc[0]['GPI']:.1f}); two archetypes, both dominated by full-service restaurants, lose money on average."
    )

    doc.add_heading("1. Introduction and problem statement", 1)
    doc.add_paragraph(
        "Not every restaurant should follow the same growth path. Some are aggregator-ready but margin fragile, some operationally strong but under-scaled, "
        "and some overextended. Conventional analytics describe what happened or what is optimal now. This project asks which restaurants are structurally "
        "positioned for sustainable growth, how they should be classified, and which levers matter most."
    )

    doc.add_heading("2. Data", 1)
    doc.add_paragraph(
        f"The dataset has {m['n_restaurants']:,} branches and 30 fields: identity (cuisine, segment, subregion), demand (orders, AOV, growth factor), "
        "channel orders, revenues and net profits for in-store, Uber Eats, DoorDash and self-delivery, cost rates, commission, and delivery logistics."
    )
    doc.add_paragraph("Data quality checks performed:")
    _bullets(
        doc,
        [
            "No missing values and no duplicate branches; RestaurantID is unique.",
            "Channel order counts sum exactly to MonthlyOrders for every row.",
            "Channel net profits reproduce Revenue x (1 - COGS - OPEX [- Commission]) and self-delivery cost equals orders x cost per order, within rounding.",
            "The supplied InStoreShare column deviates from InStoreOrders / MonthlyOrders by up to 0.20 (correlation 0.995), so shares were recomputed from order counts. The supplied UE, DD and SD shares match delivery-order shares.",
            "Column names in the file differ slightly from the brief (for example InStoreOrders instead of InStoreOrdersCount, DeliveryCostPerOrder instead of DeliveryCostOrder).",
            "GrowthFactor takes only four values and is fully determined by subregion (see Section 4). It carries location information rather than branch-level momentum.",
        ],
    )

    doc.add_heading("3. Methodology", 1)
    doc.add_heading("3.1 Feature engineering and KPIs", 2)
    _table(
        doc,
        pd.DataFrame(
            {
                "KPI / feature": ["Scale", "Cost Discipline", "Aggregator Dependence", "Expansion Headroom", "Revenue Quality", "Channel Balance", "Net Margin"],
                "Definition used": [
                    "MonthlyOrders x GrowthFactor",
                    "1 - (COGSRate + OPEXRate); cross-sectional, so level stands in for stability",
                    "(UberEats + DoorDash orders) / MonthlyOrders",
                    "DeliveryRadiusKM per 100 delivery orders",
                    "AOV x NetMargin x margin consistency, where consistency = 1 / (1 + std of the four channel margins)",
                    "Normalized Shannon entropy of the four channel order shares",
                    "Total net profit / total revenue across channels",
                ],
            }
        ),
    )
    doc.add_heading("3.2 Preprocessing", 2)
    _bullets(
        doc,
        [
            "Log transform on heavily skewed variables (scale, revenue, headroom, revenue quality, delivery cost ratio), then Yeo-Johnson power transform and standardization to remove multi-scale dominance.",
            "One-hot encoding of segment, cuisine and subregion, down-weighted to 0.35 so that categories inform but do not dominate the geometry.",
        ],
    )
    doc.add_heading("3.3 Dimensionality reduction and clustering", 2)
    doc.add_paragraph(
        f"PCA retained the components needed for 90% of variance ({m['pca_components']} components). K-Means (20-30 restarts, fixed seed), Ward hierarchical clustering "
        "and DBSCAN (eps from the 90th percentile of k-distance) were run on the PCA scores. "
        "Archetype names were assigned by matching each cluster centroid to a signed feature profile for each archetype using optimal one-to-one assignment, so labels are reproducible rather than hand-picked."
    )
    doc.add_heading("3.4 Growth Potential Index", 2)
    _table(
        doc,
        pd.DataFrame(
            {
                "Component": ["Growth signal (30%)", "Cost resilience (30%)", "Channel balance (20%)", "Logistics scalability (20%)"],
                "Inputs (percentile-ranked)": [
                    "Scale, GrowthFactor, AOV",
                    "Net margin, 1 - cost burden, 1 - commission drag, revenue quality",
                    "Channel entropy, 1 - aggregator dependence, aggregator margin",
                    "Self-delivery margin, expansion headroom, 1 - delivery cost ratio, radius",
                ],
            }
        ),
    )
    doc.add_paragraph("The weighted sum is min-max scaled to 0-100. Weights are configurable in src/config.py.")
    doc.add_heading("3.5 Recommendation rules", 2)
    _bullets(
        doc,
        [
            "Optimize: GPI in the top third.",
            "Rebalance channels: GPI below the top third and either aggregator dependence in the top third or channel balance as the weakest GPI component.",
            "Hold / Stabilize: everything else.",
        ],
    )

    doc.add_heading("4. Exploratory data analysis", 1)
    doc.add_paragraph(
        f"Channel economics are the starkest pattern. In-store net margin is {ch['InStore'][0]:.1f}% and self-delivery {ch['SelfDelivery'][0]:.1f}%, "
        f"while Uber Eats ({ch['UberEats'][0]:.1f}%) and DoorDash ({ch['DoorDash'][0]:.1f}%) earn almost nothing after a roughly 30% commission. "
        f"{ch['UberEats'][1]:.0f}% of branches lose money on Uber Eats and {ch['DoorDash'][1]:.0f}% on DoorDash, versus {ch['SelfDelivery'][1]:.0f}% on self-delivery and none in store."
    )
    _fig(doc, "01_channel_economics.png", "Figure 1. Channel margin and share of branches with a negative channel profit.")
    doc.add_paragraph(
        f"Cost structure is determined by segment. Full-service branches carry a mean COGS + OPEX rate of {seg.loc['Full-service', 'CostBurden'] * 100:.1f}% and a mean net margin of "
        f"{seg.loc['Full-service', 'NetMargin'] * 100:.1f}%, whereas ghost kitchens run at {seg.loc['Ghost Kitchen', 'CostBurden'] * 100:.1f}% with a {seg.loc['Ghost Kitchen', 'NetMargin'] * 100:.1f}% margin."
    )
    _fig(doc, "02_segment_cost_margin.png", "Figure 2. Net margin and cost burden by segment.")
    doc.add_paragraph(
        f"GrowthFactor equals {sub.loc['CBD', 'GrowthFactor']:.2f} in the CBD, {sub.loc['North Shore', 'GrowthFactor']:.2f} on the North Shore, "
        f"{sub.loc['West Auckland', 'GrowthFactor']:.2f} in West Auckland and {sub.loc['South Auckland', 'GrowthFactor']:.2f} in South Auckland for every branch, so the CBD is the only contracting area."
    )
    _fig(doc, "03_growth_by_subregion.png", "Figure 3. Growth factor is constant within each subregion.", 4.8)

    doc.add_heading("5. Results", 1)
    doc.add_heading("5.1 Latent structure (PCA)", 2)
    doc.add_paragraph(
        f"The first three components explain {sum(m['pca_variance_explained'][:3]) * 100:.1f}% of variance. PC1 ({m['pca_variance_explained'][0] * 100:.1f}%) is a cost pressure / profitability axis "
        "(cost burden against self-delivery margin, aggregator margin, net margin). PC2 is a logistics reach axis (delivery radius, headroom, delivery cost). "
        "PC3 is a channel leverage axis (self-delivery share against aggregator dependence and commission drag). Growth momentum is weak in the latent space because growth is constant within a subregion."
    )
    _fig(doc, "04_pca_variance.png", "Figure 4. PCA variance explained.", 4.8)
    _fig(doc, "05_pca_loadings.png", "Figure 5. Strongest loadings on PC1-PC3.", 4.6)

    doc.add_heading("5.2 Clustering and validation", 2)
    kk = k.set_index("k")
    doc.add_paragraph(
        f"Silhouette is highest at k = 3 ({kk.loc[3, 'silhouette']:.3f}) and declines slowly to {kk.loc[6, 'silhouette']:.3f} at k = 6. "
        "k = 6 was chosen because it is the smallest solution in which the brief's strategic archetypes are all separable (cost level x delivery radius x channel mix), "
        f"at a modest silhouette cost. K-Means and Ward clustering agree with ARI {m['robustness']['kmeans_vs_hierarchical_ARI']:.2f}. "
        f"DBSCAN found {m['robustness']['dbscan_clusters']} dense region with {m['robustness']['dbscan_noise_share'] * 100:.1f}% noise, indicating a continuous cloud rather than isolated islands; "
        "the archetypes should therefore be read as strategic segments, not natural species. A random forest recovers the cluster from the features with "
        f"{m['rf_cluster_accuracy'] * 100:.1f}% held-out accuracy."
    )
    _fig(doc, "06_k_selection.png", "Figure 6. Model selection for k.")
    _fig(doc, "07_cluster_map.png", "Figure 7. Cluster map in PCA space.", 5.6)

    doc.add_heading("5.3 Archetypes", 2)
    t = summ.copy()
    t = t.reset_index()
    t["GPI"] = t["GPI"].round(1)
    t["NetMargin"] = (t["NetMargin"] * 100).round(1).astype(str) + "%"
    t["AggregatorDependence"] = (t["AggregatorDependence"] * 100).round(0).astype(int).astype(str) + "%"
    t["CostBurden"] = (t["CostBurden"] * 100).round(0).astype(int).astype(str) + "%"
    t["DeliveryRadiusKM"] = t["DeliveryRadiusKM"].round(1)
    _table(doc, t[["ClusterLabel", "Restaurants", "GPI", "NetMargin", "CostBurden", "AggregatorDependence", "DeliveryRadiusKM"]].rename(
        columns={"ClusterLabel": "Archetype", "CostBurden": "COGS+OPEX", "AggregatorDependence": "Aggregator dep.", "DeliveryRadiusKM": "Radius km"}))
    doc.add_paragraph(
        f"Loss-making archetypes: High-Growth / High-Risk and Overextended, Low Return together contain {int(summ.loc[['High-Growth / High-Risk', 'Overextended, Low Return'], 'Restaurants'].sum())} branches, "
        f"almost all full-service, with combined monthly net profit of {(prof['High-Growth / High-Risk'] + prof['Overextended, Low Return']) / 1e6:.2f}M. "
        "They differ in reach: the High-Growth / High-Risk group has the highest order scale but a short radius, while the Overextended group pairs a long radius with the same cost burden."
    )
    _fig(doc, "08_gpi_by_cluster.png", "Figure 8. GPI distribution by archetype.", 5.6)
    _fig(doc, "09_cluster_radar.png", "Figure 9. GPI component radar by archetype.", 5.0)

    doc.add_heading("5.4 Strategic classification", 2)
    rc = m["recommendation_counts"]
    doc.add_paragraph(
        f"Optimize: {rc['Optimize']} branches; Rebalance channels: {rc['Rebalance channels']}; Hold / Stabilize: {rc['Hold / Stabilize']}. "
        "Optimize branches average a GPI of "
        f"{d.loc[d.Recommendation == 'Optimize', 'GPI'].mean():.1f} and a net margin of {d.loc[d.Recommendation == 'Optimize', 'NetMargin'].mean() * 100:.1f}%."
    )
    _fig(doc, "10_recommendation_mix.png", "Figure 10. Recommendation mix by archetype.", 5.6)

    doc.add_heading("5.5 Which levers matter most", 2)
    tl = lev.head(5)
    doc.add_paragraph(
        "Ranked by the share of variance that separates the archetypes (eta squared), the strongest levers are "
        + ", ".join(f"{n} ({v:.2f})" for n, v in tl.items())
        + ". Margin and cost levers dominate because full-service economics split the network in two; channel mix and delivery radius then separate the profitable groups."
    )
    _fig(doc, "11_lever_importance.png", "Figure 11. Lever importance.", 4.8)

    doc.add_heading("6. Strategic recommendations", 1)
    _bullets(
        doc,
        [
            "Scalable Self-Delivery Leaders and Stable Local Performers: prioritise for capital; widen delivery radius selectively and protect margin.",
            "Aggregator-Dependent Low Margin and Operationally Strong, Under-Scaled: rebalance toward in-store and self-delivery, renegotiate commission, and use marketplaces for acquisition rather than as the primary channel.",
            "High-Growth / High-Risk and Overextended, Low Return: hold expansion; fix COGS and OPEX before adding volume, because growth multiplies a negative margin.",
            "CBD: contraction (growth factor 0.99) suggests defensive, not expansion, investment.",
        ],
    )

    doc.add_heading("7. Limitations", 1)
    _bullets(
        doc,
        [
            "Single monthly snapshot: cost 'stability' and growth momentum cannot be measured over time; GrowthFactor is constant within subregion.",
            "Silhouette scores near 0.2 mean clusters overlap; treat labels as segments and use the continuous GPI for ranking.",
            "GPI weights are expert-set, not estimated against a realised outcome. Sensitivity of rankings to the weights should be tested before capital allocation.",
            "Aggregator commission is applied uniformly in the source data, so channel economics reflect those assumptions.",
        ],
    )
    doc.add_heading("8. Conclusion", 1)
    doc.add_paragraph(
        "Growth readiness at SkyCity is driven less by demand than by structure: segment-level cost burden, reliance on low-margin aggregators, and delivery economics. "
        "The GPI and archetype framework turns those structural differences into branch-level actions and is delivered as a live Streamlit dashboard."
    )
    REPORTS.mkdir(parents=True, exist_ok=True)
    out = REPORTS / "Research_Paper_Growth_Potential_Modeling.docx"
    doc.save(out)
    return out


def build_exec_summary():
    d, m, summ, lev, k, ch, seg, sub, prof = _facts()
    doc = _base()
    doc.add_heading("Executive Summary", 0)
    doc.add_paragraph("Growth Potential Modeling for SkyCity Auckland Restaurants & Bars")
    doc.add_heading("Why this matters", 1)
    doc.add_paragraph(
        "Expansion funding is more effective when it targets branches that are structurally ready to grow. "
        f"We assessed {m['n_restaurants']:,} branches on demand, cost, channel mix and delivery reach and assigned each a Growth Potential Score (0-100), an archetype and a recommended action."
    )
    doc.add_heading("Key findings", 1)
    neg = (d["TotalNetProfit"] < 0).mean() * 100
    _bullets(
        doc,
        [
            f"{neg:.0f}% of branches are loss-making overall. They are almost all full-service restaurants, whose cost base (about {seg.loc['Full-service', 'CostBurden'] * 100:.0f}% of revenue) is structurally too high.",
            f"Third-party platforms are nearly margin-free: Uber Eats and DoorDash return about {ch['UberEats'][0]:.1f}% net margin versus {ch['InStore'][0]:.0f}% in store and {ch['SelfDelivery'][0]:.0f}% on self-delivery. About {ch['UberEats'][1]:.0f}% of branches lose money on each platform.",
            f"Only {rc_opt(m)} branches ({m['recommendation_counts']['Optimize'] / m['n_restaurants'] * 100:.0f}%) are ready for investment now (Optimize); {m['recommendation_counts']['Rebalance channels']} need channel rebalancing and {m['recommendation_counts']['Hold / Stabilize']} should hold and stabilise.",
            "The Auckland CBD is the only subregion with a contracting growth factor (0.99); the other subregions grow 3-5% per month in the data.",
        ],
    )
    doc.add_heading("Six branch archetypes", 1)
    t = summ.reset_index()[["ClusterLabel", "Restaurants", "GPI", "NetMargin"]].copy()
    t["GPI"] = t["GPI"].round(1)
    t["NetMargin"] = (t["NetMargin"] * 100).round(1).astype(str) + "%"
    t["Policy stance"] = t["ClusterLabel"].map(
        {
            "Scalable Self-Delivery Leaders": "Invest and scale",
            "Stable Local Performers": "Invest selectively",
            "Aggregator-Dependent Low Margin": "Rebalance channels",
            "Operationally Strong, Under-Scaled": "Rebalance, then scale",
            "High-Growth / High-Risk": "Fix costs first",
            "Overextended, Low Return": "Stabilise or restructure",
        }
    )
    _table(doc, t.rename(columns={"ClusterLabel": "Archetype", "NetMargin": "Net margin"}))
    _fig(doc, "07_cluster_map.png", "Branches projected into two dimensions and coloured by archetype.", 5.0)
    doc.add_heading("Recommendations for stakeholders", 1)
    _bullets(
        doc,
        [
            "Direct growth capital to branches flagged Optimize, led by the Scalable Self-Delivery Leaders and Stable Local Performers.",
            "Support channel rebalancing: encourage direct and self-delivery ordering and commission renegotiation for aggregator-heavy branches.",
            "Pause expansion support for loss-making full-service branches until cost ratios are corrected.",
            "Use defensive rather than expansion measures in the CBD.",
            "Refresh the scoring monthly using time-series data to measure growth momentum and cost stability directly.",
        ],
    )
    doc.add_heading("Caveat", 1)
    doc.add_paragraph(
        "Results come from a single monthly snapshot and expert-set scoring weights. Archetypes overlap, so the continuous score should guide ranking and the archetype should guide the type of action."
    )
    out = REPORTS / "Executive_Summary_Government_Stakeholders.docx"
    doc.save(out)
    return out


def rc_opt(m):
    return m["recommendation_counts"]["Optimize"]


if __name__ == "__main__":
    print(build_paper())
    print(build_exec_summary())
